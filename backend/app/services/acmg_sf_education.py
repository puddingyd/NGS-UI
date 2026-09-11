"""Patient-facing ACMG SF education, independent of case findings and PGx.

Keep medical text and citations in the versioned JSON catalogue. This renderer
only lays out the catalogue; it never interprets a patient's variants.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CONTENT_PATH = Path(__file__).resolve().parents[1] / "report_content" / "acmg_sf_v3_3.json"


def load_catalogue() -> dict:
    return json.loads(CONTENT_PATH.read_text(encoding="utf-8"))


def _font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{kind}"), name)


def _paragraph(doc, text: str, font: str, *, size: float = 11,
               bold: bool = False, heading: int | None = None,
               before: float = 0, after: float = 4, keep_next: bool = False,
               label: str = ""):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(15 if size >= 11 else 13)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    paragraph._p.get_or_add_pPr().append(snap)
    fmt.keep_together = True
    fmt.keep_with_next = keep_next or heading is not None
    if heading is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(heading))
        paragraph._p.get_or_add_pPr().append(outline)
    if label:
        _font(paragraph.add_run(label), font, size, bold=True)
    _font(paragraph.add_run(text), font, size, bold=bold)
    return paragraph


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _link(paragraph, text: str, font: str, *, size: float = 10,
          anchor: str = "", url: str = "") -> None:
    link = OxmlElement("w:hyperlink")
    if anchor:
        link.set(qn("w:anchor"), anchor)
    else:
        link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    link.set(qn("w:history"), "1")
    run = paragraph.add_run(text)
    _font(run, font, size)
    run.underline = True
    link.append(run._r)
    paragraph._p.append(link)


def _index_table(doc, entries: list[tuple[int, dict]], font: str) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = doc.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    widths = [int(available * fraction) for fraction in (0.07, 0.31, 0.32, 0.30)]
    for column, width in zip(table.columns, widths):
        column.width = width
    props = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 45), ("bottom", 45), ("left", 90), ("right", 90)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    props.append(margins)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if edge in ("bottom", "insideH") else "nil")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    props.append(borders)
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    rows = [(None, ["編號", "相關疾病", "基因", "遺傳模式"])]
    rows += [(condition, [f"{number:02}", condition["title"],
                          "、".join(condition["genes"]), condition["index_inheritance"]])
             for number, condition in entries]
    for index, (condition, values) in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            fmt = p.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = Pt(13)
            # Keep a category heading/header with its first record on a page.
            fmt.keep_with_next = index == 0
            if condition and col == 1:
                _link(p, value, font, anchor=f"acmgsf_{condition['id']}")
            else:
                _font(p.add_run(value), font, 10, bold=index == 0)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shade)


def render_acmg_sf_education(doc, *, font_name: str = "MingLiU") -> None:
    """Append the complete index, disease introductions and bibliography."""
    catalogue = load_catalogue()
    conditions = list(enumerate(catalogue["conditions"], 1))
    # Bibliography numbering is independent of disease numbering, so merging
    # diseases does not renumber the reference list approved by the reviewer.
    source_keys = catalogue["source_order"]
    source_numbers = {key: index for index, key in enumerate(source_keys, 1)}
    existing_ids = doc.element.xpath(".//w:bookmarkStart/@w:id")
    bookmark_id = max((int(value) for value in existing_ids), default=0) + 1

    _paragraph(doc, catalogue["title"], font_name, size=13, bold=True, heading=0, after=7)
    _paragraph(doc, catalogue["introduction"], font_name, after=7)
    _paragraph(doc, "疾病索引", font_name, size=12, bold=True, heading=1)
    _paragraph(doc, catalogue["reading_note"], font_name, size=10, before=3, after=9, keep_next=True)
    for category in catalogue["categories"]:
        entries = [(number, item) for number, item in conditions if item["category"] == category["id"]]
        _paragraph(doc, category["title"], font_name, bold=True, heading=2, before=6, after=4)
        _index_table(doc, entries, font_name)

    # Break on the heading itself: a separate break paragraph can overflow a
    # full index page and create an otherwise empty page in Word/LibreOffice.
    introduction = _paragraph(doc, "疾病介紹", font_name, size=13, bold=True, heading=1, after=7)
    introduction.paragraph_format.page_break_before = True
    for category in catalogue["categories"]:
        _paragraph(doc, category["title"], font_name, size=12, bold=True, heading=2, before=9, after=7)
        for number, condition in conditions:
            if condition["category"] != category["id"]:
                continue
            heading = _paragraph(doc, f"{number:02}　{condition['title']} {condition['english']}", font_name,
                                 size=12, bold=True, heading=3, before=9, after=3)
            _bookmark(heading, f"acmgsf_{condition['id']}", bookmark_id)
            bookmark_id += 1
            fields = [("相關基因：", "、".join(condition["genes"])),
                      ("遺傳模式：", condition["inheritance"]),
                      ("可能表現與病程：", condition["clinical_course"]),
                      ("追蹤與治療：", condition["management"])]
            if condition["notes"]:
                fields.append(("補充說明：", condition["notes"]))
            for index, (label, text) in enumerate(fields):
                last = index == len(fields) - 1
                _paragraph(doc, text, font_name, label=label,
                           keep_next=not last, after=12 if last else 4)

    bibliography = _paragraph(doc, "ACMG SF 疾病簡介參考資料", font_name,
                              size=13, bold=True, heading=1, after=7)
    bibliography.paragraph_format.page_break_before = True
    _paragraph(doc, f"查閱日期：{catalogue['sources_accessed']}。電子版可點選文章名稱開啟原文。",
               font_name, size=10, after=9)
    for key in source_keys:
        source = catalogue["sources"][key]
        p = _paragraph(doc, f"[{source_numbers[key]}] ", font_name, size=10, after=2, keep_next=True)
        _bookmark(p, f"acmgsf_source_{key}", bookmark_id)
        bookmark_id += 1
        _link(p, source["title"], font_name, url=source["url"])
        _font(p.add_run(f". {source['publisher']}."), font_name, 10)
        _paragraph(doc, source["url"], font_name, size=9, after=7)
