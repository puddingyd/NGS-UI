import gzip
import json
from pathlib import Path

import pytest
from docx import Document

from app.services import gba, docx_export, manual_acmg, sample_loader, sample_layout


def payload(haplotypes=None, variants=None, cn=4):
    return {"totalCopyNumber": cn, "deletionBreakpointInGene": None,
            "recombinantHaplotypes": haplotypes if haplotypes is not None else ["RecNciI", ""],
            "variants": variants if variants is not None else []}


def rec():
    variants, replaced, warnings = gba.project(payload(), [], {}, "hg38")
    assert not warnings
    return next(iter(variants.values()))


def write_vcf(path, sample="SOURCE", records=None):
    path.write_text("##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\t" + sample + "\n" + "\n".join(records or []))


def test_recncii_is_one_heterozygous_allele_with_whole_allele_classification():
    v = rec()
    assert v["zygosity"] == "het"
    assert v["CLNSIG"] == "Pathogenic"
    assert v["clinvar_variation_id_old"] == "4297"
    assert v["depth"] is None and v["alt_af"] is None
    assert v["gba"]["deletion_breakpoint_in_gene"] is None
    assert v["gba_composition"] == "變異 RecNciI 包含位於同一等位基因上的 c.[1448T>C;1483G>C;1497G>C]（p.[Leu483Pro;Ala495Pro;Val499=]）。"
    assert manual_acmg.normalize_variant_id(v["id"]) == v["id"]
    assert manual_acmg.is_snv_variant_id(v["id"])


@pytest.mark.parametrize("name", list(gba.HAPLOTYPES) + ["1263del", "1263del+RecTL"])
def test_every_supported_named_haplotype_is_preserved(name):
    variants, _, warnings = gba.project(payload([name, ""]), [], {}, "hg38")
    assert len(variants) == 1 and not warnings
    v = next(iter(variants.values()))
    assert v["gba"]["components"]
    assert not v["gba"]["review_required"]
    assert v["zygosity"] == "het"
    assert v["gba_composition"]
    if name != "RecNciI":
        assert v.get("CLNSIG") != "Pathogenic"  # not a blanket P rule


def test_homozygous_and_two_different_recombinant_alleles():
    variants, _, _ = gba.project(payload(["RecNciI", "RecNciI"]), [], {}, "hg38")
    assert len(variants) == 1 and next(iter(variants.values()))["zygosity"] == "hom"
    variants, _, _ = gba.project(payload(["RecNciI", "L483P"]), [], {}, "hg38")
    assert len(variants) == 2
    assert all(v["zygosity"] == "het" for v in variants.values())


def test_hgvs_encoded_haplotype_uses_same_identity_as_named_caller_result():
    variants, _, _ = gba.project(payload(["NM_000157.4:c.1483G>C+NM_000157.4:c.1448T>C", ""]), [], {}, "hg38")
    assert list(variants) == [rec()["id"]]
    assert next(iter(variants.values()))["HGVS_C"] == "RecNciI"


def test_projection_keeps_conventional_second_variant_and_manual_curation(monkeypatch):
    v = rec()
    monkeypatch.setattr(manual_acmg, "bulk_current", lambda *a, **k: {})
    monkeypatch.setattr(manual_acmg, "bulk_observed_counts", lambda *a, **k: {})
    vs = {v["id"]: v}
    categories = {}
    sample_loader._apply_effective_acmg(vs, categories, sample_id="CASE", genome_build="hg38", meta={})
    assert vs[v["id"]]["effective_acmg_class"] == "Pathogenic"
    assert vs[v["id"]]["effective_acmg_source"] == "GBA1 allele curation"
    restored = docx_export._report_clinvar_variants(vs)
    assert restored[v["id"]]["CLNSIG"] == "Pathogenic"
    sample_loader._apply_effective_acmg(vs, categories, sample_id="CASE", genome_build="hg38", meta={
        "edits": {v["id"]: {"manual_acmg": {"classification": "Uncertain significance", "score": 0}}}})
    assert vs[v["id"]]["effective_acmg_class"] == "Uncertain significance"


def test_unknown_and_no_call_are_not_negative_or_pathogenic():
    variants, _, _ = gba.project(payload(["new-haplotype", ""]), [], {}, "hg38")
    v = next(iter(variants.values()))
    assert v["gba"]["review_required"] and v["CLNSIG"] != "Pathogenic"
    with pytest.raises(ValueError, match="尚待釐清"):
        gba.validate_report(variants, variants)
    variants, _, warnings = gba.project(payload([None, None], cn=None), [], {}, "hg38")
    assert not variants and warnings


def test_multiallelic_targeted_gt_and_joint_evidence(tmp_path):
    path = tmp_path / "source.vcf"
    write_vcf(path, records=[
        "chr1\t155235217\t.\tC\tG,A\t50\tPASS\tTARGETED;ALLELE_ID=.,NM_000157.4:c.1483G>C,NM_000157.4:c.1483G>T\tGT:JDP:JAD:JAF\t0/2:80:20,10,50:0.125,0.625",
        "chr1\t155235252\t.\tA\tG\t0\tTargetedWT\tTARGETED\tGT\t0/0",
        "chr1\t155214590\t.\tC\tG\t50\tPASS\tTARGETED\tGT\t0/1",
    ])
    rows = gba.read_evidence(path, "SOURCE", "hg38")
    assert len(rows) == 1
    assert rows[0]["ALT"] == "A"
    assert rows[0]["allele_id"] == "NM_000157.4:c.1483G>T"
    call = {"alleleId": rows[0]["allele_id"], "alleleCopyNumber": 1, "genotypeQuality": 25, "filter": "PASS"}
    variants, _, _ = gba.project(payload(["", ""], [call]), rows, {}, "hg38")
    v = next(iter(variants.values()))
    assert v["id"] == "chr1-155235217-C-A"
    assert v["zygosity"] == "het" and v["alt_af"] is None
    assert v["gba"]["evidence"][0]["format"]["JDP"] == "80"
    with pytest.raises(ValueError, match="sample"):
        gba.read_evidence(path, "OTHER", "hg38")


def test_targeted_vcf_without_joint_markers_uses_hard_filtered_and_keeps_other_alleles(tmp_path):
    targeted = tmp_path / "SOURCE.targeted.vcf.gz"
    hard_filtered = tmp_path / "SOURCE.hard-filtered.vcf.gz"
    def write_gzip(path, records):
        text = ("##fileformat=VCFv4.2\n"
                "#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tSOURCE\n"
                + "\n".join(records) + "\n")
        with gzip.open(path, "wt") as handle:
            handle.write(text)
    write_gzip(targeted, [
        "chr1\t155235900\t.\tA\tG\t50\tPASS\tTARGETED;ALLELE_ID=.,NM_000157.4:c.1226A>G\tGT:DP:AD:AF\t0/1:50:25,25:0.5",
    ])
    write_gzip(hard_filtered, [
        "chr1\t155235203\t.\tC\tG\t50\tPASS\tTARGETED;Recombinant\tGT:JDP:JAD:JAF\t0/1:62:16,46:0.74",
        "chr1\t155235217\t.\tC\tG\t50\tPASS\tTARGETED;Recombinant;ALLELE_ID=.,NM_000157.4:c.1483G>C\tGT:JDP:JAD:JAF\t0/1:68:16,52:0.76",
        "chr1\t155235252\t.\tA\tG\t50\tPASS\tTARGETED;Recombinant;ALLELE_ID=.,NM_000157.4:c.1448T>C\tGT:JDP:JAD:JAF\t0/1:61:15,46:0.75",
    ])
    source, rows = gba._select_evidence(hard_filtered, "SOURCE", "hg38")
    assert source == targeted
    assert len(rows) == 4
    data = payload(variants=[{"alleleId": "NM_000157.4:c.1226A>G",
                              "alleleCopyNumber": 1, "filter": "PASS"}])
    variants, _, warnings = gba.project(data, rows, {}, "hg38")
    assert not warnings and len(variants) == 2
    recncii = next(v for v in variants.values() if v["gba"]["name"] == "RecNciI")
    assert [r["format"]["JDP"] for r in recncii["gba"]["evidence"]] == ["61", "68", "62"]
    assert any(v["gba"]["name"] == "NM_000157.4:c.1226A>G" for v in variants.values())


def test_load_recovers_joint_display_evidence_from_unchanged_source(tmp_path, monkeypatch):
    source = tmp_path / "SOURCE.hard-filtered.vcf.gz"
    source.write_bytes(b"unchanged source")
    stat = source.stat()
    variant = rec()
    sidecar = tmp_path / "CASE.gba.json"
    sidecar.write_text(json.dumps({
        "schema_version": gba.SCHEMA_VERSION, "sample_id": "CASE",
        "source_sample_id": "SOURCE", "status": "complete", "genome_build": "hg38",
        "raw_gba": payload(), "variants": {variant["id"]: variant},
        "sources": {"vcf": str(source), "vcf_signature": [stat.st_size, stat.st_mtime_ns]},
    }))
    rows = [
        {"id": f"chr1-{pos}-{ref}-G", "CHROM": "chr1", "POS": pos,
         "REF": ref, "ALT": "G", "allele_id": allele_id, "copies": 1,
         "gt": "0/1", "filter": "PASS", "info": {"TARGETED": True, "Recombinant": True},
         "format": {"GT": "0/1", "JDP": dp, "JAD": ad, "JAF": af}}
        for pos, ref, allele_id, dp, ad, af in (
            (155235203, "C", "", "62", "16,46", "0.74"),
            (155235217, "C", "NM_000157.4:c.1483G>C", "68", "16,52", "0.76"),
            (155235252, "A", "NM_000157.4:c.1448T>C", "61", "15,46", "0.75"),
        )
    ]
    monkeypatch.setattr(gba.sample_layout, "state_file", lambda *_: sidecar)
    monkeypatch.setattr(gba.sample_layout, "snv_raw_tsv", lambda *_: tmp_path / "missing.tsv")
    monkeypatch.setattr(gba, "_select_evidence", lambda *_args, **_kwargs: (source, rows))
    loaded = gba.load("CASE")["variants"][variant["id"]]
    assert [r["format"]["JDP"] for r in loaded["gba"]["evidence"]] == ["61", "68", "62"]
    assert loaded["CLNSIG"] == variant["CLNSIG"]
    assert loaded["POS"] == 155235203


def test_components_are_replaced_but_unrelated_variant_is_kept(tmp_path, monkeypatch):
    path = tmp_path / "source.vcf"
    write_vcf(path, records=[
        "chr1\t155235217\t.\tC\tG\t50\tPASS\tTARGETED;Recombinant;ALLELE_ID=.,NM_000157.4:c.1483G>C\tGT:JAF\t0/1:0.75",
        "chr1\t155235252\t.\tA\tG\t50\tPASS\tTARGETED;Recombinant;ALLELE_ID=.,NM_000157.4:c.1448T>C\tGT\t0/1",
        "chr1\t155235203\t.\tC\tG\t50\tPASS\tTARGETED;Recombinant\tGT\t0/1",
    ])
    rows = gba.read_evidence(path, "SOURCE", "hg38")
    variants, replaced, _ = gba.project(payload(), rows, {}, "hg38")
    assert len(variants) == 1 and len(replaced) == 3
    monkeypatch.setattr(gba, "load", lambda _: {"variants": variants, "replaced_ids": replaced})
    current = {r["id"]: {} for r in rows}
    current["other"] = {"gene_symbol": "GBA1", "HGVS_C": "c.1226A>G"}
    gba.merge(current, "CASE")
    assert set(current) == set(variants) | {"other"}


def test_source_discovery_is_exact_and_sample_mismatch_fails(tmp_path):
    vcf = tmp_path / "vcf.gz" / "SOURCE.hard-filtered.vcf.gz"
    vcf.parent.mkdir()
    other = tmp_path / "other" / "SOURCE" / "germline_seq"
    other.mkdir(parents=True)
    json_path = other / "SOURCE.targeted.json"
    json_path.write_text(json.dumps({"sampleId": "OTHER", "genomeBuild": "hg38", "gba": payload()}))
    assert gba.discover_json(vcf, "SOURCE") == json_path
    post = tmp_path / "post"
    post.mkdir()
    with pytest.raises(ValueError, match="sampleId"):
        gba.prepare_outputs(mode="dragen", source_vcf=vcf, source_sample_id="SOURCE", sample_id="CASE", post_dir=post, raw_tsv=tmp_path / "raw")
    assert not list(post.iterdir())


def test_no_source_and_inhouse_write_explicit_unavailable(tmp_path):
    for mode in ("dragen", "inhouse"):
        result = gba.prepare_outputs(mode=mode, source_vcf=tmp_path / "SOURCE.vcf", source_sample_id="SOURCE", sample_id="CASE", post_dir=tmp_path, raw_tsv=tmp_path / "raw")
        assert result["status"] == "not_available"
        assert result["variants"] == {}
        assert (tmp_path / "CASE.gba.json").is_file()


def test_carrier_category_and_second_bullet_follow_existing_template():
    v = rec()
    variants = {v["id"]: v}
    assert docx_export._health_acmg_categorized_ids(list(variants), variants, {}) == ([], list(variants))
    doc = Document()
    docx_export._health_snv_gene_block(doc, [(v, {})], disease_text="高雪氏症（Gaucher disease）", inheritance_codes=["AR"], sex_karyotype="")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "2. " + v["gba_composition"] in text
    assert "3. 此為致病性之變異位點" in text
    assert "4. 建議至遺傳諮詢" in text
    assert "符合帶因者狀態" in text
    ref = docx_export._snv_reference_text(v, {}, acmg_zh=True)
    assert v["gba_composition"] in ref
    assert "VCV000004297.9" in ref
    assert "顯示其為罕見" not in ref


def test_recncii_health_table_missing_fields_use_aligned_fullwidth_dashes():
    v = rec()
    doc = Document()
    docx_export._health_snv_gene_block(
        doc, [(v, {})], disease_text="高雪氏症（Gaucher disease）",
        inheritance_codes=["AR"], sex_karyotype="",
    )
    table = [p.text for p in doc.paragraphs if "RS ID" in p.text][0]
    data_line = next(line for p in doc.paragraphs for line in p.text.splitlines()
                     if line.startswith("     GBA1"))
    first_dash = data_line.index("－")
    second_dash = data_line.index("－", first_dash + 1)
    assert docx_export._str_width(data_line[:first_dash]) == 5 + 9
    assert docx_export._str_width(data_line[:second_dash]) == 5 + 9 + 7
    assert docx_export._str_width(data_line) == docx_export._str_width(table.splitlines()[0])


def test_health_report_negative_limitation_is_conditional():
    assert docx_export._HEALTH_NEGATIVE_LIMITATION.startswith(
        "若於本次檢測之基因中未檢出疾病資料庫中已收錄之致病性或疑似致病性變異，並不能排除"
    )


def test_recncii_plus_separate_snv_first_category_with_phase_unknown():
    v = rec()
    snv = {"id": "chr1-155235843-T-C", "gene_symbol": "GBA1", "HGVS_C": "c.1226A>G", "HGVS_P": "p.Asn409Ser", "zygosity": "het", "ACMG_classification": "Pathogenic"}
    variants = {v["id"]: v, snv["id"]: snv}
    assert docx_export._health_acmg_categorized_ids(list(variants), variants, {}) == (list(variants), [])
    doc = Document()
    docx_export._health_snv_gene_block(doc, [(v, {}), (snv, {})], disease_text="高雪氏症（Gaucher disease）", inheritance_codes=["AR"], sex_karyotype="")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.count(v["gba_composition"]) == 1
    assert "2. " + v["gba_composition"] in text
    assert "3. 此為致病性之變異位點。由於兩變異之相位尚未確認" in text
    assert "符合帶因者狀態" not in text


def test_gba_secondary_does_not_require_fake_vaf(monkeypatch):
    v = rec()
    assert v["alt_af"] is None
    assert sample_loader._is_secondary_snv_candidate(v)
    monkeypatch.setattr(sample_loader, "_secondary_panel_genes", lambda: {"carrier": {"GBA1"}, "acmg_sf": {"BRCA1"}})
    assert sample_loader._build_secondary_snv_categories({v["id"]: v}) == {"carrier": [v["id"]], "acmg_sf": []}
    v["gba"]["review_required"] = ["FILTER 未通過"]
    assert not sample_loader._is_secondary_snv_candidate(v)


def test_old_component_selection_cannot_silently_become_new_report():
    sample = {"variants": {}, "gba": {"replaced_ids": ["chr1-155235252-A-G"]}}
    with pytest.raises(ValueError, match="整合為 allele"):
        gba.validate_saved_selection(sample, {"status": {"chr1-155235252-A-G": "1"}})
    with pytest.raises(ValueError, match="整合為 allele"):
        gba.validate_saved_selection(sample, {"secondary_findings": {"carrier": {"selected": ["chr1-155235252-A-G"]}}}, panels={"carrier"})


def test_structural_and_copy_neutral_haplotypes_do_not_share_manual_identity():
    neutral = rec()
    data = payload(cn=3)
    data["deletionBreakpointInGene"] = True
    variants, _, _ = gba.project(data, [], {}, "hg38")
    assert next(iter(variants)) != neutral["id"]
