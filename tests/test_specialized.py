import gzip
import importlib.util
import json
from pathlib import Path

import pytest
from docx import Document

from app.services import specialized as sp, gba, sample_layout, sample_loader, docx_export, manual_acmg


def payload():
    return {"sampleId": "SOURCE", "softwareVersion": "dragen v4.3.16", "genomeBuild": "hg38",
            "gba": {"totalCopyNumber": 4, "deletionBreakpointInGene": None, "recombinantHaplotypes": ["RecNciI", ""], "variants": []},
            "cyp21a2": {"totalCopyNumber": 4, "deletionBreakpointInGene": None, "recombinantHaplotypes": ["", ""], "variants": []},
            "smn": {"smn1CopyNumber": 2, "smn2CopyNumber": 1, "smn2Delta78CopyNumber": 1,
                    "totalCopyNumberFloat": "3.58", "fullLengthCopyNumberFloat": "3.42",
                    "variants": [{"hgvs": "NM_000344.4:c.*3+80T>G", "altCopyNumber": 0, "altCopyNumberQuality": 150}]},
            "hba": {"totalCopyNumber": 4, "genotype": "aa/aa", "genotypeQuality": 54, "genotypeFilter": "PASS", "variants": []}}


def record(pos=32038514, chrom="chr6", ref="C", alt="T", aid="NM_000500.9:c.92C>T", gt="0/1", source="targeted", filt="PASS", joint=False):
    fmt = "GT:JDP:JAD:JAF:JGT" if joint else "GT:DP:AD:AF"
    values = f"{gt}:80:20,60:0.75:0/0/0/1" if joint else f"{gt}:40:20,20:0.5"
    return sp._parse_record(f"{chrom}\t{pos}\t{chrom}_{pos}_{pos}\t{ref}\t{alt}\t50\t{filt}\tTARGETED;ALLELE_ID=.,{aid}\t{fmt}\t{values}", source, "hg38")


def setup_bundle(tmp_path, monkeypatch, data=None):
    source = tmp_path / "original"
    source.mkdir()
    data = data or payload()
    (source / "SOURCE.targeted.json").write_text(json.dumps(data))
    vcf = source / "SOURCE.hard-filtered.vcf.gz"
    rows = [record(155235252, "chr1", "A", "G", "NM_000157.4:c.1448T>C", joint=True),
            record(155235217, "chr1", "C", "G", "NM_000157.4:c.1483G>C", joint=True),
            record(155235203, "chr1", "C", "G", "", joint=True), record(gt="0/0")]
    with gzip.open(vcf, "wt") as f:
        f.write("##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tSOURCE\n")
        for r in rows:
            f.write(r["raw"] + "\n")
    post = tmp_path / "tertiary" / "08_postprocessing"
    post.mkdir(parents=True)
    raw = post.parent / "raw.tsv"
    raw.write_text("CHROM\tPOS\tREF\tALT\tGENE\n")
    result = sp.prepare_outputs(mode="dragen", source_vcf=vcf, source_sample_id="SOURCE", sample_id="CASE", post_dir=post, raw_tsv=raw)
    monkeypatch.setattr(sample_layout, "state_file", lambda sample, name: post / f"{sample}.{name}")
    monkeypatch.setattr(sample_layout, "snv_raw_tsv", lambda *_: raw)
    return source, post, raw, result


def test_bundle_is_self_contained_and_keeps_gba_identity(tmp_path, monkeypatch):
    source, post, raw, result = setup_bundle(tmp_path, monkeypatch)
    before = next(v for v in result["variants"].values() if v.get("gba"))
    assert before["id"] == next(iter(gba.project(payload()["gba"], [], {}, "hg38")[0]))
    # No source exists; any even attempted stat/open is a test failure.
    for path in source.iterdir():
        path.unlink()
    orig_stat, orig_open = Path.stat, Path.open
    def stat(path, *a, **k):
        assert source not in path.parents, "original path stat forbidden"
        return orig_stat(path, *a, **k)
    def opened(path, *a, **k):
        assert source not in path.parents, "original file open forbidden"
        return orig_open(path, *a, **k)
    monkeypatch.setattr(Path, "stat", stat)
    monkeypatch.setattr(Path, "open", opened)
    after = gba.load("CASE")["variants"][before["id"]]
    assert after["gba"]["evidence"][0]["format"]["JDP"] == "80"
    assert sp.load("CASE")["loci"]["smn"]["smn1_copy_number"] == 2
    stage = tmp_path / "offline"
    stage.mkdir()
    rebuilt = sp.rebuild_local(sample_id="CASE", post_dir=stage, raw_tsv=raw)
    assert set(rebuilt["variants"]) == set(result["variants"])
    assert sum(p.stat().st_size for p in post.iterdir()) < 100_000
    doc = Document()
    docx_export._snv_gene_block(doc, [(after, {})], tier="第一類")
    assert "RecNciI" in "\n".join(p.text for p in doc.paragraphs)


def test_reference_records_and_header_preserved_not_positive(tmp_path, monkeypatch):
    _, post, _, result = setup_bundle(tmp_path, monkeypatch)
    with gzip.open(post / "CASE.specialized.evidence.jsonl.gz", "rt") as f:
        evidence = [json.loads(x) for x in f]
    assert any(x["type"] == "header" for x in evidence)
    assert any(x.get("format", {}).get("GT") == "0/0" for x in evidence)
    assert not result["loci"]["cyp21a2"]["finding_ids"]
    assert not result["loci"]["smn"]["finding_ids"]
    assert not result["loci"]["hba"]["finding_ids"]


def test_checksum_and_sample_mismatch_fail_closed(tmp_path, monkeypatch):
    _, post, _, _ = setup_bundle(tmp_path, monkeypatch)
    (post / "CASE.specialized.calls.json").write_text("{}")
    with pytest.raises(ValueError, match="checksum"):
        sp.load("CASE")


def test_json_and_vcf_samples_never_mixed(tmp_path):
    src = tmp_path / "SOURCE.vcf"
    src.write_text("##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tOTHER\n")
    (tmp_path / "SOURCE.targeted.json").write_text(json.dumps(payload()))
    with pytest.raises(ValueError, match="sample"):
        sp.prepare_outputs(mode="dragen", source_vcf=src, source_sample_id="SOURCE", sample_id="CASE", post_dir=tmp_path / "post", raw_tsv=tmp_path / "raw")
    data = payload(); data["sampleId"] = "OTHER"
    (tmp_path / "SOURCE.targeted.json").write_text(json.dumps(data))
    with pytest.raises(ValueError, match="sampleId"):
        sp.prepare_outputs(mode="dragen", source_vcf=src, source_sample_id="SOURCE", sample_id="CASE", post_dir=tmp_path / "post", raw_tsv=tmp_path / "raw")


def test_size_limit_and_large_unindexed_vcf(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "MAX_BYTES", 64)
    p = tmp_path / "large.json"; p.write_bytes(b"x" * 64)
    with pytest.raises(ValueError, match="100 MB"):
        sp._read_small(p)
    with pytest.raises(ValueError, match="100 MB"):
        sp._write_small(tmp_path / "out", b"x" * 64)
    with pytest.raises(ValueError, match="索引"):
        sp._extract(p, "hard_filtered", "SOURCE", "hg38")
    assert not (tmp_path / "out").exists()


def test_cyp_haplotype_components_count_once_and_phase_unknown_second_allele():
    data = payload()
    aid1, aid2 = "NM_000500.9:c.710T>A", "NM_000500.9:c.713T>A"
    data["cyp21a2"]["recombinantHaplotypes"] = [aid1 + "+" + aid2, ""]
    rows = [record(aid=aid1), record(pos=32038517, aid=aid2)]
    result = sp.project(data, rows, {}, "hg38")
    vs = {k:v for k,v in result["variants"].items() if v.get("specialized")}
    assert len(vs) == 1
    v = next(iter(vs.values()))
    assert v["specialized"]["phase"] == "caller_haplotype"
    assert v["zygosity"] == "het" and len(v["specialized"]["components"]) == 2
    assert v["CLNSIG"] != "Pathogenic"
    assert len(result["loci"]["cyp21a2"]["replaced_ids"]) == 2
    assert docx_export._health_acmg_categorized_ids(list(vs), vs, {}) == ([], list(vs))
    second = {**v, "id": "second", "specialized": {**v["specialized"], "kind": "variant", "phase": "unknown"}}
    vs["second"] = second
    assert docx_export._health_acmg_categorized_ids(list(vs), vs, {})[1] == []


@pytest.mark.parametrize("cn,expected", [(0,"risk"), (1,"carrier"), (2,None), (3,None), (None,None)])
def test_smn_copy_number_not_fake_heterozygous_snv(cn, expected):
    data = payload(); data["smn"]["smn1CopyNumber"] = cn
    r = sp.project(data, [], {}, "hg38")
    vs = {k:v for k,v in r["variants"].items() if v.get("specialized")}
    if expected:
        v = next(iter(vs.values()))
        assert v["zygosity"] == f"CN = {cn}"
        assert v["POS"] is None and v["REF"] == "" and v["ALT"] == ""
        assert v["specialized"]["health_category"] == expected
        risk, carriers = docx_export._health_acmg_categorized_ids(list(vs), vs, {})
        assert bool(carriers) == (expected == "carrier")
        assert v["CLNSIG"] != "Pathogenic"
        assert sample_loader._is_secondary_snv_candidate(v)
    else:
        assert not vs
    if cn is None:
        assert r["loci"]["smn"]["status"] == "review_required"


def test_smn_marker_is_not_confirmed_carrier_or_reportable_pathogenic():
    data = payload(); data["smn"]["variants"][0]["altCopyNumber"] = 1
    result = sp.project(data, [], {}, "hg38")
    vs = {k:v for k,v in result["variants"].items() if v.get("specialized")}
    v = next(iter(vs.values()))
    assert v["specialized"]["kind"] == "risk_marker"
    assert not sample_loader._is_secondary_snv_candidate(v)
    with pytest.raises(ValueError, match="尚待釐清"):
        sp.validate_selection({"variants": vs}, {"status": {v["id"]:"1"}})


@pytest.mark.parametrize("gt,cn,category", [("-a3.7/aa",3,"carrier"),("--/aa",2,"carrier"),("-a3.7/-a4.2",2,"carrier"),
    ("--/-a3.7",1,"risk"),("--/--",0,"risk"),("aaa3.7/aa",5,None),("--/aaa3.7",3,"carrier"),("--/aaa4.2",3,"carrier"),("unknown/aa",4,None)])
def test_hba_genotype_preserved_and_categorized(gt, cn, category):
    data = payload(); data["hba"].update(genotype=gt,totalCopyNumber=cn)
    result = sp.project(data, [], {}, "hg38")
    v = next(v for v in result["variants"].values() if v.get("specialized"))
    assert v["specialized"]["genotype"] == gt
    assert v["specialized"]["health_category"] == category
    assert "SEA" not in v["HGVS_C"]
    assert v["zygosity"] not in {"het", "hom"}
    assert manual_acmg.is_snv_variant_id(v["id"])
    if category:
        risk, carriers = docx_export._health_acmg_categorized_ids([v["id"]], {v["id"]:v}, {})
        assert bool(carriers) == (category == "carrier")
    if "unknown" in gt:
        assert v["specialized"]["hard_blocked"]


def test_hba_joint_small_variant_not_assigned_by_guess():
    data = payload(); aid="NM_000517.6:c.427T>C"
    data["hba"]["variants"]=[{"alleleId":aid,"alleleCopyNumber":1,"filter":"PASS","genotypeQuality":50}]
    result = sp.project(data, [record(173000,"chr16",aid=aid,joint=True)], {}, "hg38")
    v = next(v for v in result["variants"].values() if v.get("specialized"))
    assert v["zygosity"] == ""
    assert v["specialized"]["review_required"]
    assert not v["CLNSIG"]


def test_multiallelic_correct_index_and_ref_alt_conflict_not_hidden():
    r = sp._parse_record("chr6\t32038514\t.\tC\tT,A\t50\tPASS\tALLELE_ID=.,NM_000500.9:c.92C>T,NM_000500.9:c.92C>A\tGT:JDP:JAD:JAF\t0/2:80:20,10,50:0.125,0.625", "targeted", "hg38")
    a = sp.allele_evidence([r])[0]
    assert a["allele_id"] == "NM_000500.9:c.92C>A" and a["ALT"] == "A" and a["allele_index"] == 2
    data=payload(); data["cyp21a2"]["variants"]=[{"alleleId":"NM_000500.9:c.92C>T","alleleCopyNumber":1,"filter":"PASS"}]
    rows=[record(),record(gt="0/0",source="hard_filtered")]
    result=sp.project(data,rows,{},"hg38")
    v=next(v for v in result["variants"].values() if v.get("specialized"))
    assert any("不一致" in x for x in v["specialized"]["review_required"])


def test_review_confirmation_requires_current_evidence_and_comment():
    data=payload(); data["cyp21a2"]["variants"]=[{"alleleId":"NM_000500.9:c.92C>T","alleleCopyNumber":1,"filter":"PASS"}]
    result=sp.project(data,[record(joint=True)],{},"hg38")
    v=next(v for v in result["variants"].values() if v.get("specialized")); vid=v["id"]
    report={"status":{vid:"1"},"edits":{vid:{"specialized_review":v["specialized"]["review_token"],"comment":"獨立驗證已確認 CYP21A2 定位"}}}
    sp.validate_selection(result, report)
    report["edits"][vid]["specialized_review"]="stale"
    with pytest.raises(ValueError,match="尚待釐清"):
        sp.validate_selection(result,report)


def test_health_dosage_requires_manual_classification_and_keeps_template():
    data=payload();data["smn"]["smn1CopyNumber"]=1
    r=sp.project(data,[],{},"hg38");v=next(v for v in r["variants"].values() if v.get("specialized"));vid=v["id"]
    report={"secondary_findings":{"carrier":{"selected":[vid]}}}
    with pytest.raises(ValueError,match="人工分類"):
        sp.validate_selection(r,report,panels={"carrier"})
    report["edits"]={vid:{"manual_acmg":{"classification":"Pathogenic"}}}
    sp.validate_selection(r,report,panels={"carrier"})
    doc=Document()
    docx_export._health_snv_gene_block(doc,[(v,report["edits"][vid])],disease_text=sp.DISEASES["SMN1"],inheritance_codes=["AR"],sex_karyotype="XX")
    text="\n".join(p.text for p in doc.paragraphs)
    assert "CN = 1" in text and "帶因者" in text and "2. SMN1" in text
    assert "Heterozygous" not in text
    assert "罕見" not in docx_export._snv_reference_text(v,{}) or "不據此推定罕見" in docx_export._snv_reference_text(v,{})


def test_publish_rolls_back_every_file(tmp_path,monkeypatch):
    spec=importlib.util.spec_from_file_location("rebuild_specialized",Path(__file__).parents[1]/"scripts/rebuild_specialized.py")
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    stage=tmp_path/"stage";dest=tmp_path/"dest";stage.mkdir();dest.mkdir()
    for name in sp.MANAGED_NAMES:
        (stage/f"CASE.{name}").write_text("new")
        (dest/f"CASE.{name}").write_text("old")
    original=module.os.replace
    def fail(src,dst):
        if Path(src).parent==stage and Path(src).name.endswith("manifest.json"):
            raise OSError("injected failure")
        return original(src,dst)
    monkeypatch.setattr(module.os,"replace",fail)
    with pytest.raises(OSError):module.publish(stage,dest,"CASE")
    assert all(p.read_text()=="old" for p in dest.iterdir())


def test_health_method_scope_changes_only_when_targeted_results_reported():
    doc = Document()
    docx_export._section_methods(doc, "WGS", health=True, specialized_genes="SMN1、HBA1/HBA2")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "SMN1、HBA1/HBA2 的專一性分析結果" in text
    assert "除上述專一性分析外" in text
    assert "不代表本檢測已完整涵蓋其他基因" in text


def test_hba_complex_health_group_requires_consistent_current_manual_category():
    data = payload()
    data["hba"].update(genotype="--/aa", totalCopyNumber=2,
        variants=[{"alleleId": "NM_000517.6:c.427T>C", "alleleCopyNumber": 1, "filter": "PASS"}])
    result = sp.project(data, [record(173000, "chr16", aid="NM_000517.6:c.427T>C", joint=True)], {}, "hg38")
    variants = {k: v for k, v in result["variants"].items() if v.get("specialized")}
    report = {"secondary_findings": {"carrier": {"selected": list(variants)}}, "edits": {}}
    for vid, v in variants.items():
        report["edits"][vid] = {"manual_acmg": {"classification": "Pathogenic"},
            "specialized_review": v["specialized"]["review_token"], "comment": "獨立驗證與 locus 判讀依據"}
    with pytest.raises(ValueError, match="指定健檢"):
        sp.validate_selection(result, report, panels={"carrier"})
    for vid, v in variants.items():
        report["edits"][vid].update(specialized_health_category="risk", specialized_health_token=v["specialized"]["review_token"])
    sp.validate_selection(result, report, panels={"carrier"})
    risk, carriers = docx_export._health_acmg_categorized_ids(list(variants), variants, report["edits"])
    assert set(risk) == set(variants) and not carriers
    report["edits"][vid]["specialized_health_category"] = "carrier"
    with pytest.raises(ValueError, match="分類不一致"):
        sp.validate_selection(result, report, panels={"carrier"})


def test_stale_selection_missing_locus_and_implicit_selection_validation():
    with pytest.raises(ValueError, match="不在目前資料包"):
        sp.validate_selection({"variants": {}}, {"status": {"SMN-old": "1"}})
    r = sp.project({}, [], {}, "hg38")
    assert all(v["status"] == "not_available" for v in r["loci"].values())
    notes = sp.report_limitations(r, {"SMN1"})
    assert len(notes) == 1 and "不能據此視為陰性" in notes[0]
    data = payload(); data["smn"]["smn1CopyNumber"] = 1
    r = sp.project(data, [], {}, "hg38")
    vid = next(k for k, v in r["variants"].items() if v.get("specialized"))
    with pytest.raises(ValueError, match="人工分類"):
        sp.validate_selection(r, {}, panels={"carrier"}, selected_ids=[vid])


def test_bundle_moves_without_original_paths(tmp_path, monkeypatch):
    import shutil
    source, post, raw, result = setup_bundle(tmp_path, monkeypatch)
    moved = tmp_path / "moved"
    shutil.move(str(post), moved)
    shutil.rmtree(source)
    monkeypatch.setattr(sample_layout, "state_file", lambda sample, name: moved / f"{sample}.{name}")
    assert sp.load("CASE")["variants"] == result["variants"]


def test_cyp_partial_components_stay_reviewable_and_cannot_double_report():
    data = payload()
    data["cyp21a2"]["recombinantHaplotypes"] = ["NM_000500.9:c.710T>A+NM_000500.9:c.713T>A", ""]
    result = sp.project(data, [record(aid="NM_000500.9:c.710T>A")], {}, "hg38")
    v = next(v for v in result["variants"].values() if v.get("specialized"))
    assert "缺少完整 PASS allele 證據" in v["specialized"]["review_required"]
    component_id = v["specialized"]["evidence"][0]["id"]
    with pytest.raises(ValueError, match="同時回報"):
        sp.validate_selection(result, {"status": {v["id"]: "1", component_id: "1"}})


def test_hba_reachable_in_real_carrier_panel_and_inherited_manual_acmg(monkeypatch):
    data = payload(); data["hba"].update(genotype="--/aa", totalCopyNumber=2)
    result = sp.project(data, [], {}, "hg38")
    v = next(v for v in result["variants"].values() if v.get("specialized")); vid = v["id"]
    assert {"HBA1", "HBA2"} <= sample_loader._secondary_panel_genes()["carrier"]
    variants = {vid: v}; tiers = {}
    monkeypatch.setattr(manual_acmg, "bulk_current", lambda *a, **k: {})
    monkeypatch.setattr(manual_acmg, "bulk_observed_counts", lambda *a, **k: {})
    sample_loader._apply_effective_acmg(variants, tiers, sample_id="CASE", genome_build="hg38", meta={
        "edits": {vid: {"manual_acmg": {"classification": "Pathogenic", "score": None}}}})
    assert variants[vid]["effective_acmg_class"] == "Pathogenic"
    cats = sample_loader._build_secondary_snv_categories(variants)
    assert vid in cats["carrier"] and vid not in cats["acmg_sf"]
