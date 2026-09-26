from docx import Document

from app.services import docx_export


def _gjb2_variant(
    variant_id: str,
    *,
    rs_id: str,
    hgvs_c: str,
    hgvs_p: str,
    acmg: str,
) -> dict:
    return {
        "id": variant_id,
        "gene_symbol": "GJB2",
        "ensembl_transcript": "ENST00000382848",
        "refseq_transcript": "NM_004004.6",
        "rs_id": rs_id,
        "exon": "2/2",
        "HGVS_C": hgvs_c,
        "HGVS_P": hgvs_p,
        "zygosity": "Heterozygous",
        "CLNSIG": "Pathogenic",
        "ACMG_classification": acmg,
        "Disease1": "Deafness, autosomal recessive 1A (220290)(AR)",
    }


def test_diagnosis_groups_same_gene_snvs_and_combines_acmg_wording():
    doc = Document()
    first = _gjb2_variant(
        "v1",
        rs_id="rs80338943",
        hgvs_c="c.235del",
        hgvs_p="p.Leu79CysfsTer3",
        acmg="Pathogenic",
    )
    second = _gjb2_variant(
        "v2",
        rs_id="rs72474224",
        hgvs_c="c.109G>A",
        hgvs_p="p.Val37Ile",
        acmg="Likely pathogenic",
    )

    docx_export._section_results(
        doc,
        {"variants": {"v1": first, "v2": second}},
        {"status": {"v1": "1", "v2": "1"}, "edits": {}},
        "WES",
    )

    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert text.count("GJB2 (ENST00000382848; NM_004004.6)") == 1
    assert text.count("ACMG&AMP指引") == 1
    assert "c.235del" in text
    assert "c.109G>A" in text
    assert text.count("GJB2為Deafness, autosomal recessive 1A的致病基因之一") == 1
    assert "此為致病性及疑似致病性之變異位點，與臨床症狀相關。" in text


def test_diagnosis_empty_first_category_uses_reviewed_wording():
    doc = Document()

    docx_export._section_results(
        doc,
        {"patient_phenotype": [{"label": "Developmental delay"}]},
        {"status": {}, "edits": {}},
        "WES",
    )

    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "未找到與臨床症狀相關基因之已知致病性變異位點。" in text
    assert "在非特定" not in text


def test_em_dash_uses_full_width_for_ascii_table_padding():
    assert docx_export._str_width("—") == 2
    assert docx_export._pad_right("—", 13) == "—" + (" " * 11)


def test_cnv_report_disease_is_selected_union_plus_free_text():
    edits = {
        "report_disease_items": {
            "omim:GENE:1:1": {
                "label": "Disease A", "source": "omim", "gene": "GENE",
                "phenotype_mim": "600001", "inheritance": "AD",
            },
            "overlap:p_loss:Disease B": {
                "label": "Disease B", "source": "overlap", "overlap_type": "p_loss",
            },
            "duplicate": {"label": "disease a", "source": "overlap"},
        },
        "disease": "Disease A、Manual disease",
    }

    assert docx_export._cnv_report_disease(edits) == (
        "Disease A、Disease B、Manual disease"
    )


def test_single_gene_cnv_keeps_one_point_and_uses_combined_disease_text():
    doc = Document()
    variant = {
        "id": "cnv1", "source": "cnv", "CHROM": "17", "POS": 100,
        "END": 200, "sv_type": "DEL", "copy_number": 1, "zygosity": "het",
        "acmg_class": 5,
        "genes": [{
            "gene": "COL1A1", "location": "exon1-exon2", "omim_id": "120150",
            "omim_phenotype": "Legacy disease (600000)(AD)",
            "omim_inheritance": "AD",
        }],
    }
    edits = {
        "report_disease_items": {
            "omim:COL1A1:120150:1": {"label": "Disease A", "source": "omim"},
            "overlap:p_loss:Disease B": {"label": "Disease B", "source": "overlap"},
        },
        "disease": "Manual disease",
    }

    docx_export._cnv_variant_block(
        doc, variant, tier="1", is_wgs=True, edits=edits
    )

    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "COL1A1為Disease A、Disease B、Manual disease的致病基因之一" in text
    assert "Phenotype MIM number" not in text


def test_single_overlap_disease_does_not_inherit_gene_omim_metadata():
    doc = Document()
    variant = {
        "id": "cnv1", "source": "cnv", "CHROM": "17", "POS": 100,
        "END": 200, "sv_type": "DEL", "copy_number": 1, "zygosity": "het",
        "acmg_class": 5,
        "genes": [{
            "gene": "COL1A1", "location": "exon1-exon2", "omim_id": "120150",
            "omim_phenotype": "Legacy disease (600000)(AD)",
            "omim_inheritance": "AD",
        }],
    }
    edits = {
        "report_disease_items": {
            "overlap:p_loss:Overlap disease": {
                "label": "Overlap disease", "source": "overlap",
            },
        },
    }

    docx_export._cnv_variant_block(
        doc, variant, tier="1", is_wgs=True, edits=edits
    )

    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "COL1A1為Overlap disease的致病基因之一" in text
    assert "Phenotype MIM number" not in text
    assert "顯性遺傳" not in text
