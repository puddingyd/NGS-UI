import csv
import gzip
import json
import os
import sqlite3
import time
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from starlette.middleware.sessions import SessionMiddleware

from app import config
from app.auth import current_user
from app.services import somatic, sample_layout, sample_loader
from app.workers.somatic_run import Worker, Cancelled, subtract_vcf, vep_rows, write_rows, prepare_allele_filters
from app.routers import somatic as router


@pytest.fixture
def setup(tmp_path, monkeypatch):
    sid = "S1"
    state = tmp_path / "state"
    state.mkdir()
    raw = tmp_path / "raw.tsv"
    raw.write_text("CHROM\tPOS\tREF\tALT\tCALLERS\tZYGOSITY\tDP\tVAF\nchr1\t100\tA\tG\tDV\thet\t5\t0.1\n")
    monkeypatch.setattr(config, "DATA_ROOT", tmp_path / "data")
    monkeypatch.setattr(config, "PIPELINE_OUT_ROOT", tmp_path / "output")
    monkeypatch.setattr(sample_layout, "state_file", lambda s, name, **kw: state / f"{s}.{name}")
    monkeypatch.setattr(sample_layout, "state_dir", lambda s: state)
    monkeypatch.setattr(sample_layout, "snv_raw_tsv", lambda s: raw)
    monkeypatch.setattr(somatic, "canonical_gene", lambda s: {"OLD": "GENE1"}.get(s.upper(), s.upper()))
    somatic.atomic_json(state / f"{sid}.sample_metadata.json", {"genome_build": "hg38"})
    ref = tmp_path / "hg38.fa"
    ref.write_text(">chr1\n" + "A" * 1000 + "\n")
    Path(str(ref) + ".fai").write_text("chr1\t1000\t6\t1000\t1001\nchrX\t1000\t0\t0\t0\n")
    regions = tmp_path / "genes.bed"
    regions.write_text("#assembly=GRCh38\nchr1\t99\t110\tGENE1\texon\nchr1\t109\t120\tGENE2\texon\nchr1\t49\t300\tGENE1\tgene\n")
    cfg = {"assembly": "GRCh38", "reference": str(ref), "gene_regions": str(regions), "gene_regions_release": "test"}
    return sid, raw, cfg


def publish(sid, raw, rows, *, run_id="a" * 32, archived=False, created=1):
    target = somatic.result_dir(sid, run_id)
    target.mkdir(parents=True)
    write_rows(target / "annotations.tsv", rows)
    somatic.atomic_json(target / "candidates.json", [{"id": f'{r["CHROM"]}-{r["POS"]}-{r["REF"]}-{r["ALT"]}'} for r in rows])
    data = somatic.manifest(sid)
    data["runs"].append({"run_id": run_id, "created": created, "raw_signature": somatic.signature(raw), "archived": archived})
    somatic.atomic_json(somatic.index_path(sid), data)
    return run_id


def row(pos="101", alt="T", filt="PASS", dp="8", gene="GENE1", tx="ENST1"):
    return {"CHROM": "chr1", "POS": pos, "REF": "A", "ALT": alt, "GENE": gene,
            "TRANSCRIPT": tx, "CONSEQUENCE": "missense_variant", "SOMATIC_FILTER": filt,
            "CALLERS": "Mutect2", "DP": dp, "AD": "7,1", "VAF": "0.125", "ACMG_CRITERIA": ""}


def test_targets_union_aliases_padding_and_alleles(setup):
    _, _, cfg = setup
    result = somatic.resolve_targets({"genes": "OLD, GENE2\nGENE1", "positions": "chr1:125-150; 1:150:A>G"}, cfg)
    assert result["genes"] == ["GENE1", "GENE2"]
    assert result["intervals"] == [["chr1", 80, 150]]
    assert result["total_bases"] == 71
    assert result["positions"][-1]["alt"] == "G"
    body = somatic.resolve_targets({"genes": "GENE1", "region_mode": "gene"}, cfg)
    assert body["intervals"] == [["chr1", 50, 300]]


@pytest.mark.parametrize("payload", [{}, {"genes": "NOPE"}, {"positions": "chr1:0"},
    {"positions": "chr1:1001"}, {"positions": "chrM:1"}, {"positions": "chr1:20-10"},
    {"positions": "chr1:1:A>A"}, {"positions": "chr1:1;$(id)"}, {"genes": []}, {"region_mode": "all"}])
def test_invalid_targets_rejected(setup, payload):
    with pytest.raises(ValueError):
        somatic.resolve_targets(payload, setup[2])


def test_subtraction_exact_allele_and_outside_region(tmp_path):
    source = tmp_path / "calls.vcf"
    header = "##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\n"
    source.write_text(header + "chr1\t100\t.\tA\tG\t.\tPASS\t.\nchr1\t100\t.\tA\tT\t.\tPASS\t.\n"
                      "chr1\t101\t.\tAA\tA\t.\tweak_evidence\t.\nchr1\t900\t.\tA\tT\t.\tPASS\t.\n")
    with sqlite3.connect(":memory:") as db:
        db.execute("CREATE TABLE variants (id TEXT PRIMARY KEY)")
        db.execute("INSERT INTO variants VALUES ('chr1-100-A-G')")
        counts = subtract_vcf(source, tmp_path / "novel.vcf", db, [["chr1", 100, 102]])
    assert counts == {"germline_excluded": 1, "outside_targets": 1, "new_candidates": 2, "pass": 1}
    assert "\tA\tG\t" not in (tmp_path / "novel.vcf").read_text()
    assert "weak_evidence" in (tmp_path / "novel.vcf").read_text()


def test_somatic_low_depth_loads_and_filtered_requires_selection(setup):
    sid, raw, _ = setup
    assert somatic.summary(sid)["completed"] is False
    run = publish(sid, raw, [row(), row("102", filt="weak_evidence")])
    variants = somatic.load_variants(sid)
    assert set(variants) == {"chr1-101-A-T"}
    assert variants["chr1-101-A-T"]["depth"] == 8
    assert variants["chr1-101-A-T"]["alt_af"] == 0.125
    assert variants["chr1-101-A-T"]["somatic"] is True
    somatic.select_filtered(sid, run, "chr1-102-A-T")
    assert somatic.load_variants(sid)["chr1-102-A-T"]["somatic_filter"] == "weak_evidence"
    assert somatic.summary(sid)["completed"] is True
    with pytest.raises(ValueError):
        somatic.select_filtered(sid, run, "chr1-999-A-T")


def test_zero_result_run_still_enables_checkbox(setup):
    sid, raw, _ = setup
    publish(sid, raw, [])
    assert somatic.summary(sid)["completed"] is True
    assert somatic.load_variants(sid) == {}


def test_rerun_archives_but_preserves_marked_evidence_and_search(setup):
    sid, raw, _ = setup
    publish(sid, raw, [row()], archived=True)
    publish(sid, raw, [row("102", gene="GENE2")], run_id="b" * 32, created=2)
    assert set(somatic.load_variants(sid)) == {"chr1-102-A-T"}
    kept = somatic.load_variants(sid, keep_ids={"chr1-101-A-T"})
    assert kept["chr1-101-A-T"]["somatic_historical"] is True
    assert set(somatic.load_variants(sid, genes={"GENE2"})) == {"chr1-102-A-T"}
    assert "chr1-101-A-T" in somatic.load_variants(sid, wanted={"chr1-101-A-T"})


def test_changed_germline_requires_reconciliation(setup):
    sid, raw, _ = setup
    publish(sid, raw, [row()])
    raw.write_text(raw.read_text() + "chr1\t101\tA\tT\tDV\thet\t10\t0.1\n")
    assert somatic.summary(sid)["stale"] is True
    assert somatic.load_variants(sid) == {}


def test_missing_reviewed_history_remains_detectable_after_fresh_run(setup):
    sid, raw, _ = setup
    publish(sid, raw, [row()], archived=True)
    raw.write_text(raw.read_text() + "chr1\t110\tA\tC\tDV\thet\t30\t0.5\n")
    publish(sid, raw, [row("102")], run_id="b" * 32, created=2)
    assert somatic.summary(sid)["stale"] is False
    available = set(somatic.load_variants(sid, keep_ids={"chr1-101-A-T"}))
    assert somatic.missing_review_ids(sid, {"chr1-101-A-T", "germline-only"}, available) == ["chr1-101-A-T"]
    # A retained germline card for the same allele satisfies the review.
    assert somatic.missing_review_ids(sid, {"chr1-101-A-T"}, available | {"chr1-101-A-T"}) == []


def test_report_blocks_missing_reviewed_somatic(monkeypatch):
    from app.services import docx_export
    sample = {"somatic": {"review_missing_ids": ["chr1-101-A-T"]}}
    monkeypatch.setattr(docx_export.sample_loader, "load_sample", lambda *a, **kw: sample)
    monkeypatch.setattr(docx_export, "_report_clinvar_sample", lambda s: s)
    monkeypatch.setattr(docx_export.report_store, "load", lambda s: {})
    with pytest.raises(ValueError, match="已標記 Somatic"):
        docx_export.build_diagnosis_docx("S1")


def test_vep_preserves_mutect_evidence_and_all_transcripts(tmp_path):
    path = tmp_path / "vep.json"
    obj = {"input": "chr1\t101\t.\tA\tT\t.\tPASS\tTLOD=12\tGT:AD:AF:DP\t0/1:97,3:0.03:100",
           "transcript_consequences": [{"gene_symbol": "GENE1", "transcript_id": "ENST1", "hgvsc": "ENST1:c.1A>T", "consequence_terms": ["missense_variant"]},
                                       {"gene_symbol": "GENE2", "transcript_id": "ENST2", "consequence_terms": ["intron_variant"]}]}
    path.write_text(json.dumps(obj) + "\n")
    rows = vep_rows(path)
    assert len(rows) == 2
    assert rows[0]["AD"] == "97,3" and rows[0]["VAF"] == "0.03"
    assert rows[0]["ZYGOSITY"] == "" and rows[0]["SOMATIC_GT"] == "0/1"
    assert rows[0]["HGVS_C"] == "c.1A>T"


def test_authenticated_scoped_api(setup, monkeypatch):
    sid, raw, cfg = setup
    app = FastAPI()
    app.add_middleware(SessionMiddleware, secret_key="test-only")
    app.include_router(router.router)
    client = TestClient(app)
    assert client.get(f"/api/samples/{sid}/somatic").status_code == 401
    app.dependency_overrides[current_user] = lambda: {"username": "tester"}
    monkeypatch.setattr(somatic, "settings", lambda: cfg)
    monkeypatch.setattr(somatic, "sample_bams", lambda s: [])
    assert client.post(f"/api/samples/{sid}/somatic/preview", json={"genes": "GENE1"}).status_code == 200
    run = "c" * 32
    somatic.atomic_json(somatic.job_dir(run) / "state.json", {"sample_id": "OTHER", "status": "completed"})
    assert client.get(f"/api/samples/{sid}/somatic/jobs/{run}").status_code == 404
    assert client.post(f"/api/samples/{sid}/somatic/jobs/{run}/cancel", json={}).status_code == 404


def test_failed_worker_and_cancel_do_not_publish(setup, monkeypatch):
    sid, raw, cfg = setup
    run = "d" * 32
    jobdir = somatic.job_dir(run)
    somatic.atomic_json(jobdir / "state.json", {"sample_id": sid, "run_id": run})
    somatic.atomic_json(jobdir / "config.json", cfg)
    worker = Worker(run)
    (jobdir / "cancel").touch()
    with pytest.raises(Cancelled):
        worker.run(["not-a-command"])
    assert somatic.summary(sid)["completed"] is False


def test_somatic_is_never_secondary_finding():
    assert not sample_loader._is_secondary_snv_candidate({"somatic": True, "tier": "1A", "alt_af": 0.5})


def test_report_source_and_vaf_for_somatic_only():
    from docx import Document
    from app.services.docx_export import _snv_gene_block
    doc = Document()
    variant = {"id": "chr1-101-A-T", "gene_symbol": "GENE1", "somatic": True, "alt_af": .03,
               "somatic_filter": "PASS", "somatic_clinvar_release": "2026-07-20"}
    _snv_gene_block(doc, [(variant, {"somatic_validation": "已驗證"})], tier="1")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Somatic pipeline (Mutect2)" in text and "3.0%" in text and "已驗證" in text
    doc = Document()
    _snv_gene_block(doc, [(dict(variant, somatic=False), {})], tier="1")
    assert "Somatic pipeline" not in "\n".join(p.text for p in doc.paragraphs)


def test_worker_complete_chain_only_publishes_novel_alleles(setup, tmp_path, monkeypatch):
    """Exercise orchestration with synthetic tool outputs, not a biological validation."""
    sid, raw, cfg = setup
    bam = tmp_path / "S1.bam"
    bam.write_text("synthetic")
    cfg.update({key: str(raw) for key in ("germline_resource", "clinvar_vcf")})
    cfg.update(vep_cache=str(tmp_path), vep_cache_version="115", clinvar_release="2026-07-20")
    cfg.update({key + "_command": [key] for key in ("gatk", "samtools", "bcftools", "vep")})
    monkeypatch.setattr(config, "GENEBE_DB", tmp_path / "missing.genebe.gz")
    run = "e" * 32
    targets = somatic.resolve_targets({"genes": "GENE1"}, cfg)
    record = {"sample_id": sid, "run_id": run, "created": 1, "targets": targets,
              "raw_signature": somatic.signature(raw), "bam_signature": somatic.signature(bam), "bam_path": str(bam)}
    somatic.atomic_json(somatic.job_dir(run) / "state.json", record)
    somatic.atomic_json(somatic.job_dir(run) / "config.json", cfg)
    commands = []
    header = "##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tS1\n"
    vcf = header + "".join(f"chr1\t{p}\t.\tA\t{a}\t.\t{f}\tTLOD=10\tGT:AD:AF:DP\t0/1:97,3:0.03:100\n"
                          for p, a, f in [(100, "G", "PASS"), (101, "T", "PASS"), (102, "T", "weak_evidence")])
    def fake_run(self, args, output=None):
        args = list(map(str, args))
        commands.append(args)
        self.check_cancel()
        if output:
            text = "version1\n"
            if args[1:3] == ["view", "-H"]:
                text = "@SQ\tSN:chr1\tLN:1000\n@SQ\tSN:chrX\tLN:1000\n@RG\tID:1\tSM:S1\n"
            elif args[1] == "norm":
                text = Path(args[-1]).read_text()
            elif args[1] == "depth":
                text = "chr1\t100\t100\nchr1\t101\t100\nchr1\t102\t100\n"
            output.write_text(text)
        elif args[0] == "gatk" and "-O" in args:
            dest = Path(args[args.index("-O") + 1])
            text = vcf if args[1] in {"Mutect2", "FilterMutectCalls"} else "test"
            if dest.suffix == ".gz":
                with gzip.open(dest, "wt") as handle:
                    handle.write(text)
            else:
                dest.write_text(text)
        elif args[0] == "vep":
            source = Path(args[args.index("--input_file") + 1]).read_text()
            assert "\t100\t" not in source, "Existing low-DP germline must be removed before annotation"
            dest = Path(args[args.index("--output_file") + 1])
            dest.write_text("".join(json.dumps({"input": line,
                "transcript_consequences": [{"gene_symbol": "GENE1", "transcript_id": "ENST1", "consequence_terms": ["missense_variant"]}]}) + "\n"
                for line in source.splitlines() if not line.startswith("#")))
    monkeypatch.setattr(Worker, "run", fake_run)
    Worker(run).execute()
    assert somatic.read_job(run)["status"] == "completed"
    assert somatic.read_job(run)["counts"]["germline_excluded"] == 1
    assert set(somatic.load_variants(sid)) == {"chr1-101-A-T"}
    assert raw.read_text().endswith("\t5\t0.1\n")
    assert "chr1\t100" not in (somatic.result_dir(sid, run) / "novel.vcf").read_text()
    assert len(somatic.read_json(somatic.result_dir(sid, run) / "candidates.json")) == 2
    assert any("FilterMutectCalls" in cmd for cmd in commands)
    assert any("--clinvar" in cmd for cmd in commands)


def test_submission_cannot_accept_arbitrary_bam(setup, monkeypatch):
    sid, _, cfg = setup
    monkeypatch.setattr(somatic, "settings", lambda: cfg)
    monkeypatch.setattr(somatic, "sample_bams", lambda s: [{"path": "/approved/S1.bam"}])
    with pytest.raises(ValueError, match="有效 BAM"):
        somatic.start(sid, {"genes": "GENE1", "bam_path": "/etc/passwd"})


def test_allele_specific_filter_cannot_be_promoted_to_pass(tmp_path):
    vcf = tmp_path / "multi.vcf"
    vcf.write_text("##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\n"
                   "chr1\t100\t.\tA\tG,T\t.\tPASS\tAS_FilterStatus=SITE|weak_evidence,strand_bias\n")
    prepared = tmp_path / "prepared.vcf"
    prepare_allele_filters(vcf, prepared)
    assert "NGS_UI_AS_FILTER=SITE,weak_evidence%2Cstrand_bias" in prepared.read_text()
    # Simulate Number=A selection by bcftools norm for the second ALT.
    normalized = tmp_path / "norm.vcf"
    normalized.write_text("#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\n"
                          "chr1\t100\t.\tA\tT\t.\tPASS\tNGS_UI_AS_FILTER=weak_evidence%2Cstrand_bias\n")
    with sqlite3.connect(":memory:") as db:
        db.execute("CREATE TABLE variants (id TEXT PRIMARY KEY)")
        counts = subtract_vcf(normalized, tmp_path / "novel.vcf", db, [["chr1", 100, 100]])
    assert counts["pass"] == 0
    assert "strand_bias;weak_evidence" in (tmp_path / "novel.vcf").read_text()


def test_requested_indel_survives_left_alignment_outside_interval(tmp_path):
    source = tmp_path / "norm.vcf"
    source.write_text("#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\n"
                      "chr1\t90\t.\tAA\tA\t.\tPASS\t.\n"
                      "chr1\t91\t.\tAA\tA\t.\tPASS\tNGS_UI_TARGET=1\n")
    with sqlite3.connect(":memory:") as db:
        db.execute("CREATE TABLE variants (id TEXT PRIMARY KEY)")
        counts = subtract_vcf(source, tmp_path / "novel.vcf", db, [["chr1", 100, 101]], {"chr1-90-AA-A"})
    assert counts["new_candidates"] == 2


def test_failed_publish_leaves_previous_index(setup, monkeypatch):
    sid, raw, cfg = setup
    publish(sid, raw, [row()])
    before = somatic.index_path(sid).read_bytes()
    from app.workers import somatic_run
    run = "b" * 32
    somatic.atomic_json(somatic.job_dir(run) / "state.json", {"run_id": run, "sample_id": sid})
    somatic.atomic_json(somatic.job_dir(run) / "config.json", cfg)
    monkeypatch.setattr(somatic_run, "_acquire_sample_locks", lambda *a: [])
    def fail(self):
        raise RuntimeError("synthetic annotation failure")
    monkeypatch.setattr(Worker, "execute", fail)
    # main installs termination handlers; restore them after this in-process test.
    import signal
    old = {s: signal.getsignal(s) for s in (signal.SIGTERM, signal.SIGINT)}
    try:
        assert somatic_run.main(run) == 1
    finally:
        for s, handler in old.items():
            signal.signal(s, handler)
    assert somatic.index_path(sid).read_bytes() == before
    assert somatic.read_job(run)["status"] == "failed"
    assert not (somatic.job_dir(run) / "staging").exists()
