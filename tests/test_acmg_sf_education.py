import io
from itertools import combinations
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.services import acmg_sf_education, docx_export


def test_catalogue_covers_the_report_panel_and_cites_every_disease():
    catalogue = acmg_sf_education.load_catalogue()
    panel = Path(__file__).resolve().parents[1] / "phenotype_data/custom_panels/ACMG_SF_v3.3.txt"
    panel_genes = {line.strip() for line in panel.read_text().splitlines()
                   if line.strip() and not line.startswith("#")}
    conditions = catalogue["conditions"]
    genes = {gene for condition in conditions for gene in condition["genes"]}
    assert len(conditions) == 36
    assert len(genes) == 84
    assert genes == panel_genes
    assert genes == {gene for group in docx_export._ACMG_SF_GROUPS for gene in group["genes"]}
    assert len({c["id"] for c in conditions}) == len(conditions)
    category_ids = [category["id"] for category in catalogue["categories"]]
    assert list(dict.fromkeys(c["category"] for c in conditions)) == category_ids
    for condition in conditions:
        assert re.fullmatch(r"[a-z][a-z0-9_]*", condition["id"])
        assert condition["genes"] and len(set(condition["genes"])) == len(condition["genes"])
        for field in ("title", "english", "inheritance", "index_inheritance", "clinical_course", "management"):
            assert condition[field].strip()
        assert condition["references"]
        for key in condition["references"]:
            source = catalogue["sources"][key]
            assert source["title"] and source["publisher"]
            assert source["url"].startswith("https://")
    assert len(catalogue["source_order"]) == len(set(catalogue["source_order"])) == 45
    assert set(catalogue["source_order"]) == set(catalogue["sources"])


def test_education_index_has_working_links_and_no_patient_result_labels():
    doc = Document()
    normal_xml = doc.styles["Normal"].element.xml
    acmg_sf_education.render_acmg_sf_education(doc)
    # Round-trip through a DOCX package, including its hyperlink relationships.
    buffer = io.BytesIO()
    doc.save(buffer)
    doc = Document(io.BytesIO(buffer.getvalue()))
    assert doc.styles["Normal"].element.xml == normal_xml
    catalogue = acmg_sf_education.load_catalogue()
    assert len(doc.tables) == len(catalogue["categories"])
    assert sum(len(table.rows) - 1 for table in doc.tables) == 36
    assert len(doc.element.xpath(".//w:tblHeader")) == 6
    for table in doc.tables:
        assert [cell.text for cell in table.rows[0].cells] == ["編號", "相關疾病", "基因", "遺傳模式"]
        assert all(row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for row in table.rows)
    bookmarks = doc.element.xpath(".//w:bookmarkStart/@w:name")
    assert len(bookmarks) == len(set(bookmarks))
    assert set(doc.element.xpath(".//w:hyperlink/@w:anchor")) <= set(bookmarks)
    assert {f"acmgsf_{c['id']}" for c in catalogue["conditions"]} <= set(bookmarks)
    external_urls = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
    assert {source["url"] for source in catalogue["sources"].values()} == external_urls
    text = "\n".join(doc.element.xpath(".//w:t/text()"))
    for status in ("檢出", "未檢出", "Positive", "Negative", "No Result"):
        assert status not in text
    for disease in catalogue["conditions"]:
        assert disease["clinical_course"] in text
        assert disease["management"] in text


def test_reviewed_index_and_anchored_comments_are_applied():
    doc = Document()
    acmg_sf_education.render_acmg_sf_education(doc)
    catalogue = acmg_sf_education.load_catalogue()
    by_id = {c["id"]: c for c in catalogue["conditions"]}
    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(doc.element.xpath(".//w:t/text()"))
    assert paragraphs[0] == "ACMG疾病風險基因與相關疾病簡介"
    assert paragraphs[paragraphs.index("疾病索引") + 1] == catalogue["reading_note"]
    for removed in ("如何閱讀遺傳模式", "內容更新：", "疾病短文", "資料來源：",
                    "文中編號對應", "息肉", "半顯性", "p.Cys282Tyr", "少數僅有一份 CASQ2"):
        assert removed not in text
    assert paragraphs.count("其他疾病") == 2
    assert "參考資料" in paragraphs
    assert "ACMG SF 疾病簡介參考資料" not in text
    assert "查閱日期：" not in text
    assert next(p for p in doc.paragraphs if p.text == "疾病介紹").paragraph_format.page_break_before
    bibliography = next(p for p in doc.paragraphs if p.text == "參考資料")
    assert not bibliography.paragraph_format.page_break_before
    assert not bibliography.paragraph_format.keep_with_next
    assert bibliography.paragraph_format.space_after.pt == 0
    assert not bibliography._p.xpath("./w:pPr/w:outlineLvl")
    assert not doc.element.xpath(".//w:br[@w:type='page']")
    assert {c["id"] for c in catalogue["conditions"] if c["notes"]} == {"fh"}
    assert sum(p.startswith("補充說明：") for p in paragraphs) == 1
    assert "瘜肉相關問題與癌症風險需分別考慮" in by_id["pjs"]["management"]
    assert not by_id["pjs"]["notes"]

    merged = by_id["polyposis"]
    assert set(merged["genes"]) == {"APC", "MUTYH", "BMPR1A", "SMAD4"}
    assert not {"apc", "mutyh", "jps"} & set(by_id)
    assert "APC、BMPR1A、SMAD4：體染色體顯性遺傳" in merged["inheritance"]
    assert "MUTYH：體染色體隱性遺傳" in merged["inheritance"]
    assert "出現深褐色的小斑點" in by_id["pjs"]["clinical_course"]
    assert "女性也可能在成年後出現雙腿僵硬、走路困難，或難以控制排尿、排便等症狀" in by_id["ald"]["clinical_course"]
    assert "腫瘤或其他異常變化" in by_id["tsc"]["clinical_course"]
    assert "比周圍膚色淺的斑塊" in by_id["tsc"]["clinical_course"]
    assert "病灶" not in by_id["tsc"]["clinical_course"]
    assert "皮膚色素較淡" not in by_id["tsc"]["clinical_course"]
    assert "視網膜感受光線的功能受影響" in by_id["rpe65"]["clinical_course"]
    assert "利用光線" not in by_id["rpe65"]["clinical_course"]
    assert by_id["nf2"]["title"] == "神經纖維瘤症候群第二型"
    assert by_id["nf2"]["english"] == "NF2-related schwannomatosis"
    assert by_id["fh"]["index_inheritance"] == "體染色體顯性、體染色體隱性"
    assert by_id["pgl"]["index_inheritance"] == "體染色體顯性"
    for cid in ("ald", "fabry", "otc"):
        assert by_id[cid]["index_inheritance"] == by_id[cid]["inheritance"] == "X 染色體性聯遺傳"

    index_rows = [row for table in doc.tables for row in table.rows[1:]]
    for number, (condition, row) in enumerate(zip(catalogue["conditions"], index_rows), 1):
        assert row.cells[0].text == f"{number:02}"
        assert "".join(row.cells[1]._tc.xpath(".//w:t/text()")) == condition["title"]
        assert row.cells[2].text == "、".join(condition["genes"])
        assert row.cells[3].text == condition["index_inheritance"]
        # Chinese and English share one heading and one 12-point bold run.
        heading = next(p for p in doc.paragraphs
                       if p._p.xpath("./w:bookmarkStart/@w:name") == [f"acmgsf_{condition['id']}"])
        assert heading.text == f"{number:02}　{condition['title']} {condition['english']}"
        assert all(run.font.size.pt == 12 and run.bold for run in heading.runs)
        assert not heading._p.xpath(".//w:br")
        assert condition["english"] not in paragraphs
        # A blank line follows each complete disease, including category changes.
        last_label = "補充說明：" if condition["notes"] else "追蹤與治療："
        last_value = condition["notes"] or condition["management"]
        last_index = paragraphs.index(last_label + last_value)
        if number < len(catalogue["conditions"]):
            spacer = doc.paragraphs[last_index + 1]
            assert not spacer.text
            assert spacer.paragraph_format.line_spacing.pt == 15
            assert spacer.paragraph_format.keep_with_next
        else:
            assert paragraphs[last_index + 1] == ""
            assert doc.paragraphs[last_index + 1].paragraph_format.line_spacing.pt == 15
            assert paragraphs[last_index + 2] == "參考資料"
    # Finished exports contain neither unresolved comments nor revision markup.
    assert not doc.element.xpath(".//w:ins|.//w:del|.//w:rPrChange|.//w:pPrChange|.//w:commentRangeStart")


@pytest.mark.parametrize("sections", [None] + [
    list(selected)
    for count in range(1, 5)
    for selected in combinations(("acmg_sf", "stroke", "carrier", "pgx"), count)
])
@pytest.mark.parametrize("with_finding", [False, True])
def test_health_export_selects_and_orders_education(monkeypatch, tmp_path, sections, with_finding):
    from app.services import sample_layout

    selected = set(sections if sections is not None else ["acmg_sf", "pgx"])
    disease_selected = selected.intersection(docx_export._HEALTH_DISEASE_SECTIONS)
    fixture_variants = {
        key: {"id": key, "gene_symbol": gene, "HGVS_C": f"c.{index}01A>G",
              "ACMG_classification": "Pathogenic", "Zygosity": "Heterozygous"}
        for index, (key, gene) in enumerate(
            (("acmg_sf", "LDLR"), ("stroke", "NOTCH3"), ("carrier", "CFTR")), start=1)
    }
    variants = fixture_variants if with_finding else {}
    categories = {key: [key] if with_finding else [] for key in fixture_variants}
    loaded = []

    def load_secondary(*args, **kwargs):
        loaded.append("disease")
        return {"variants": variants, "categories": categories}

    monkeypatch.setattr(docx_export.sample_loader, "load_sample", lambda *args, **kwargs: {"meta": {"Test": "WGS"}})
    monkeypatch.setattr(docx_export.sample_loader, "load_sample_secondary_snv", load_secondary)
    monkeypatch.setattr(docx_export.report_store, "load", lambda *args: {
        "edits": {}, "secondary_findings": {key: {"selected": ids} for key, ids in categories.items()},
    })
    monkeypatch.setattr(sample_layout, "state_dir", lambda *args: tmp_path)
    monkeypatch.setattr(docx_export.phenotype_scorer, "genes_for_key", lambda *args, **kwargs: {"genes": ["LDLR"]})
    groups = [{"drug": "Clopidogrel", "genes": {"CYP2C19": {"phenotype": "Poor Metabolizer"}},
               "recommendations": [{"source": "CPIC", "level": "Strong", "recommendation": "Use an alternative antiplatelet agent."}]}]
    def load_pgx(*args):
        loaded.append("pgx")
        return {"pgx": {"report_view": {"drug_groups": groups, "health_genotype_rows": [
            {"test": "CYP2C19", "allele1": "*2", "allele2": "*2"},
        ]}}}

    # Render the actual main text, methods, lists, references and appendices,
    # with populated PGx data even when its checkbox is not selected.
    monkeypatch.setattr(docx_export.sample_loader, "load_sample_pgx", load_pgx)
    payload = docx_export.build_health_docx("education-test", sections=sections)
    doc = Document(io.BytesIO(payload))
    text = "\n".join(doc.element.xpath(".//w:t/text()"))
    title = acmg_sf_education.load_catalogue()["title"]
    assert ("disease" in loaded) == bool(disease_selected)
    assert ("pgx" in loaded) == ("pgx" in selected)
    assert (title in text) == ("acmg_sf" in selected)
    assert (docx_export._HEALTH_ACMG_CAUTION in text) == ("acmg_sf" in selected)
    assert (docx_export._HEALTH_ACMG_GENE_LIST_TITLE in text) == ("acmg_sf" in selected)
    for marker in ("藥物基因體學", "官方用藥資訊查詢", "完整用藥建議",
                   "CYP2C19", "CYP2D6", "某些藥物基因", "Clopidogrel"):
        assert (marker in text) == ("pgx" in selected), marker
    external_urls = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
    for _label, url in docx_export._HEALTH_PGX_RESOURCES:
        assert (url in external_urls) == ("pgx" in selected)
    for key, variant in fixture_variants.items():
        assert (variant["HGVS_C"] in text) == (with_finding and key in selected)
    assert ("變異位點參考資料" in text) == (with_finding and bool(disease_selected))
    if "acmg_sf" in selected:
        assert len(doc.tables) == 6
        assert "此處列出清單中之84個基因" in text
        assert "36　遺傳性轉甲狀腺素蛋白類澱粉沉積症" in text
        if with_finding:
            assert text.index("變異位點參考資料") < text.index(title)
        else:
            assert "變異位點參考資料" not in text
        if "pgx" in selected:
            assert text.index(title) < text.index("\n參考資料\n") < text.index("完整用藥建議")
            assert "Use an alternative antiplatelet agent." in text
    if "pgx" in selected and disease_selected and with_finding and "acmg_sf" not in selected:
        assert text.index("變異位點參考資料") < text.index("完整用藥建議")
    if not with_finding and "acmg_sf" not in selected and "pgx" not in selected:
        assert "附錄" not in text


@pytest.mark.parametrize("sections", [[], [""], ["unknown"], ["pgx", "unknown"]])
def test_health_export_rejects_empty_or_unsupported_selection(monkeypatch, sections):
    def unexpected_load(*args, **kwargs):
        pytest.fail("Invalid selections must be rejected before loading patient data")

    monkeypatch.setattr(docx_export.sample_loader, "load_sample", unexpected_load)
    with pytest.raises(ValueError):
        docx_export.build_health_docx("selection-test", sections=sections)


@pytest.mark.parametrize("test_type", ["WES", "WGS", "TITAN-WGS"])
def test_health_methods_without_pgx_keep_continuous_numbering(test_type):
    doc = Document()
    docx_export._section_methods(doc, test_type, health=True, include_pgx=False)
    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(paragraphs)
    assert "藥物基因" not in text
    assert "CYP2D6" not in text
    numbers = [int(match.group(1)) for p in paragraphs if (match := re.match(r"\s+(\d+)\.", p))]
    assert numbers == list(range(1, len(numbers) + 1))
    assert "無法檢測出拷貝數變異" in text
    assert ("短讀長全基因體定序" in text) == (test_type != "WES")
